"""
Multi-format output generation system for Document Organizer.
Handles export to Markdown, DOCX, CSV, and JSON formats with full traceability.
"""

import csv
import json
import logging
import os
from datetime import datetime
from pathlib import Path
from typing import Dict, List, Any, Optional

try:
    from docx import Document
    from docx.shared import Inches
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    DOCX_AVAILABLE = True
except ImportError:
    DOCX_AVAILABLE = False

from .processor import ConsolidatedData
from .models import MedicalRecord, LifeHistoryRecord, CodeReviewRecord
from .azure_gpt import GPTResponse

logger = logging.getLogger(__name__)


class OutputError(Exception):
    """Raised when output generation fails."""
    pass


class OutputWriter:
    """Handles multi-format output generation."""
    
    def __init__(self, output_dir: str = "output"):
        """
        Initialize the output writer.
        
        Args:
            output_dir: Base output directory
        """
        self.output_dir = Path(output_dir)
        self.timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        
        # Create output subdirectories
        self.md_dir = self.output_dir / "md"
        self.docx_dir = self.output_dir / "docx"
        self.csv_dir = self.output_dir / "csv"
        self.json_dir = self.output_dir / "json"
        
        for dir_path in [self.md_dir, self.docx_dir, self.csv_dir, self.json_dir]:
            dir_path.mkdir(parents=True, exist_ok=True)
        
        logger.info(f"Initialized output writer with base directory: {self.output_dir}")
    
    def write_all_formats(self, consolidated_data: ConsolidatedData, 
                         mode: str, outline_data: Dict[str, Any],
                         raw_responses: List[GPTResponse]) -> Dict[str, List[str]]:
        """
        Write output in all supported formats.
        
        Args:
            consolidated_data: Consolidated and deduplicated data
            mode: Processing mode ('medical', 'life', or 'code')
            outline_data: Document collection outline
            raw_responses: Raw GPT responses for JSON backup
            
        Returns:
            Dictionary mapping format to list of generated file paths
        """
        generated_files = {
            "markdown": [],
            "docx": [],
            "csv": [],
            "json": []
        }
        
        try:
            # Generate outline first
            outline_files = self.write_outline(outline_data, mode)
            generated_files["markdown"].extend(outline_files)
            
            # Generate main summary files
            if mode == 'medical':
                md_files = self.write_medical_markdown(consolidated_data)
                csv_files = self.write_medical_csv(consolidated_data)
                generated_files["markdown"].extend(md_files)
                generated_files["csv"].extend(csv_files)
                
                if DOCX_AVAILABLE:
                    docx_files = self.write_medical_docx(consolidated_data)
                    generated_files["docx"].extend(docx_files)
                
            elif mode == 'life':
                md_files = self.write_life_markdown(consolidated_data)
                csv_files = self.write_life_csv(consolidated_data)
                generated_files["markdown"].extend(md_files)
                generated_files["csv"].extend(csv_files)
                
                if DOCX_AVAILABLE:
                    docx_files = self.write_life_docx(consolidated_data)
                    generated_files["docx"].extend(docx_files)
            
            elif mode == 'code':
                md_files = self.write_code_markdown(consolidated_data)
                csv_files = self.write_code_csv(consolidated_data)
                generated_files["markdown"].extend(md_files)
                generated_files["csv"].extend(csv_files)
                
                if DOCX_AVAILABLE:
                    docx_files = self.write_code_docx(consolidated_data)
                    generated_files["docx"].extend(docx_files)
            
            # Backup raw JSON responses
            json_files = self.backup_json_responses(raw_responses, mode)
            generated_files["json"].extend(json_files)
            
            # Write processing metadata
            metadata_files = self.write_processing_metadata(consolidated_data, mode)
            generated_files["json"].extend(metadata_files)
            
            logger.info(f"Generated output files: {sum(len(files) for files in generated_files.values())} total")
            return generated_files
            
        except Exception as e:
            logger.error(f"Failed to write output files: {e}")
            raise OutputError(f"Output generation failed: {e}")
    
    def write_outline(self, outline_data: Dict[str, Any], mode: str) -> List[str]:
        """Write document collection outline."""
        outline_file = self.md_dir / f"outline_{mode}_{self.timestamp}.md"
        
        try:
            with open(outline_file, 'w', encoding='utf-8') as f:
                outline = outline_data.get('outline', {})
                title = outline.get('title', 'Document Collection Overview')
                
                f.write(f"# {title}\n\n")
                f.write(f"**Generated:** {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}\n")
                f.write(f"**Mode:** {mode.title()}\n")
                
                # Key themes
                themes = outline_data.get('key_themes', [])
                if themes:
                    f.write(f"\n## Key Themes\n\n")
                    for theme in themes:
                        f.write(f"- {theme}\n")
                
                # Date range and document count
                date_range = outline_data.get('date_range', 'Unknown')
                doc_count = outline_data.get('total_documents', 0)
                f.write(f"\n## Collection Summary\n\n")
                f.write(f"- **Date Range:** {date_range}\n")
                f.write(f"- **Total Documents:** {doc_count}\n")
                
                # Sections
                sections = outline.get('sections', [])
                if sections:
                    f.write(f"\n## Content Organization\n\n")
                    for section in sections:
                        section_title = section.get('section_title', 'Untitled Section')
                        f.write(f"### {section_title}\n\n")
                        
                        subsections = section.get('subsections', [])
                        for subsection in subsections:
                            sub_title = subsection.get('title', 'Untitled')
                            content = subsection.get('content', 'No description')
                            doc_count = subsection.get('document_count', 0)
                            
                            f.write(f"#### {sub_title}\n")
                            f.write(f"{content}\n")
                            f.write(f"*Documents: {doc_count}*\n\n")
            
            logger.info(f"Generated outline: {outline_file}")
            return [str(outline_file)]
            
        except Exception as e:
            logger.error(f"Failed to write outline: {e}")
            return []
    
    def write_medical_markdown(self, data: ConsolidatedData) -> List[str]:
        """Write medical data to Markdown format."""
        md_file = self.md_dir / f"medical_summary_{self.timestamp}.md"
        
        try:
            with open(md_file, 'w', encoding='utf-8') as f:
                f.write("# Medical Records Summary\n\n")
                f.write(f"**Generated:** {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}\n")
                f.write(f"**Total Records:** {len(data.medical_records)}\n")
                f.write(f"**Source Files:** {len(data.source_files)}\n\n")
                
                # Deduplication summary
                if data.deduplication_log:
                    f.write("## Deduplication Summary\n\n")
                    for log_entry in data.deduplication_log:
                        f.write(f"- {log_entry}\n")
                    f.write("\n")
                
                # Medical records
                f.write("## Medical Records\n\n")
                for i, record in enumerate(data.medical_records, 1):
                    f.write(f"### {i}. {record.condition}\n\n")
                    
                    # Visit dates
                    if record.visit_dates:
                        f.write("**Visit Dates:**\n")
                        for date in sorted(record.visit_dates):
                            f.write(f"- {date.strftime('%Y-%m-%d')}\n")
                        f.write("\n")
                    
                    # Medications
                    if record.medications:
                        f.write("**Medications:**\n")
                        for med in record.medications:
                            f.write(f"- {med}\n")
                        f.write("\n")
                    
                    # Physician
                    if record.physician:
                        f.write(f"**Physician:** {record.physician}\n\n")
                    
                    # Notes
                    if record.notes:
                        f.write(f"**Notes:** {record.notes}\n\n")
                    
                    # Source files
                    f.write("**Source Files:**\n")
                    for source in record.source_files:
                        f.write(f"- {source}\n")
                    
                    f.write(f"\n**Confidence:** {record.confidence_score:.2f}\n\n")
                    f.write("---\n\n")
            
            logger.info(f"Generated medical markdown: {md_file}")
            return [str(md_file)]
            
        except Exception as e:
            logger.error(f"Failed to write medical markdown: {e}")
            return []
    
    def write_life_markdown(self, data: ConsolidatedData) -> List[str]:
        """Write life history data to Markdown format."""
        md_file = self.md_dir / f"life_history_{self.timestamp}.md"
        
        try:
            with open(md_file, 'w', encoding='utf-8') as f:
                f.write("# Life History Summary\n\n")
                f.write(f"**Generated:** {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}\n")
                f.write(f"**Total Records:** {len(data.life_history_records)}\n")
                f.write(f"**Source Files:** {len(data.source_files)}\n\n")
                
                # Deduplication summary
                if data.deduplication_log:
                    f.write("## Deduplication Summary\n\n")
                    for log_entry in data.deduplication_log:
                        f.write(f"- {log_entry}\n")
                    f.write("\n")
                
                # Sort records by start date
                sorted_records = sorted(data.life_history_records, 
                                      key=lambda r: r.start_date or datetime.min)
                
                # Life history records
                f.write("## Career Timeline\n\n")
                for i, record in enumerate(sorted_records, 1):
                    f.write(f"### {i}. {record.role}\n\n")
                    
                    # Dates
                    start_str = record.start_date.strftime('%Y-%m-%d') if record.start_date else 'Unknown'
                    end_str = record.end_date.strftime('%Y-%m-%d') if record.end_date else 'Ongoing'
                    f.write(f"**Period:** {start_str} to {end_str}\n\n")
                    
                    # Location
                    if record.location:
                        f.write(f"**Location:** {record.location}\n\n")
                    
                    # Achievements
                    if record.key_achievements:
                        f.write("**Key Achievements:**\n")
                        for achievement in record.key_achievements:
                            f.write(f"- {achievement}\n")
                        f.write("\n")
                    
                    # Source files
                    f.write("**Source Files:**\n")
                    for source in record.source_files:
                        f.write(f"- {source}\n")
                    
                    f.write(f"\n**Confidence:** {record.confidence_score:.2f}\n\n")
                    f.write("---\n\n")
            
            logger.info(f"Generated life history markdown: {md_file}")
            return [str(md_file)]
            
        except Exception as e:
            logger.error(f"Failed to write life history markdown: {e}")
            return []
    
    def write_medical_csv(self, data: ConsolidatedData) -> List[str]:
        """Write medical data to CSV format."""
        csv_file = self.csv_dir / f"medical_records_{self.timestamp}.csv"
        
        try:
            with open(csv_file, 'w', newline='', encoding='utf-8') as f:
                writer = csv.writer(f)
                
                # Header
                writer.writerow([
                    'Condition', 'Visit_Dates', 'Medications', 'Physician', 
                    'Notes', 'Source_Files', 'Confidence_Score'
                ])
                
                # Data rows
                for record in data.medical_records:
                    visit_dates_str = '; '.join([d.strftime('%Y-%m-%d') for d in record.visit_dates])
                    medications_str = '; '.join(record.medications)
                    sources_str = '; '.join(record.source_files)
                    
                    writer.writerow([
                        record.condition,
                        visit_dates_str,
                        medications_str,
                        record.physician,
                        record.notes,
                        sources_str,
                        f"{record.confidence_score:.3f}"
                    ])
            
            logger.info(f"Generated medical CSV: {csv_file}")
            return [str(csv_file)]
            
        except Exception as e:
            logger.error(f"Failed to write medical CSV: {e}")
            return []
    
    def write_life_csv(self, data: ConsolidatedData) -> List[str]:
        """Write life history data to CSV format."""
        csv_file = self.csv_dir / f"life_history_{self.timestamp}.csv"
        
        try:
            with open(csv_file, 'w', newline='', encoding='utf-8') as f:
                writer = csv.writer(f)
                
                # Header
                writer.writerow([
                    'Role', 'Start_Date', 'End_Date', 'Location', 
                    'Key_Achievements', 'Source_Files', 'Confidence_Score'
                ])
                
                # Data rows
                for record in data.life_history_records:
                    start_date_str = record.start_date.strftime('%Y-%m-%d') if record.start_date else ''
                    end_date_str = record.end_date.strftime('%Y-%m-%d') if record.end_date else 'Ongoing'
                    achievements_str = '; '.join(record.key_achievements)
                    sources_str = '; '.join(record.source_files)
                    
                    writer.writerow([
                        record.role,
                        start_date_str,
                        end_date_str,
                        record.location,
                        achievements_str,
                        sources_str,
                        f"{record.confidence_score:.3f}"
                    ])
            
            logger.info(f"Generated life history CSV: {csv_file}")
            return [str(csv_file)]
            
        except Exception as e:
            logger.error(f"Failed to write life history CSV: {e}")
            return []
    
    def write_medical_docx(self, data: ConsolidatedData) -> List[str]:
        """Write medical data to DOCX format."""
        if not DOCX_AVAILABLE:
            logger.warning("python-docx not available, skipping DOCX generation")
            return []
        
        docx_file = self.docx_dir / f"medical_summary_{self.timestamp}.docx"
        
        try:
            doc = Document()
            
            # Title
            title = doc.add_heading('Medical Records Summary', 0)
            title.alignment = WD_ALIGN_PARAGRAPH.CENTER
            
            # Metadata
            doc.add_paragraph(f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
            doc.add_paragraph(f"Total Records: {len(data.medical_records)}")
            doc.add_paragraph(f"Source Files: {len(data.source_files)}")
            
            # Records
            for i, record in enumerate(data.medical_records, 1):
                doc.add_heading(f"{i}. {record.condition}", level=1)
                
                if record.visit_dates:
                    p = doc.add_paragraph()
                    p.add_run("Visit Dates: ").bold = True
                    dates_str = ', '.join([d.strftime('%Y-%m-%d') for d in record.visit_dates])
                    p.add_run(dates_str)
                
                if record.medications:
                    p = doc.add_paragraph()
                    p.add_run("Medications: ").bold = True
                    p.add_run(', '.join(record.medications))
                
                if record.physician:
                    p = doc.add_paragraph()
                    p.add_run("Physician: ").bold = True
                    p.add_run(record.physician)
                
                if record.notes:
                    p = doc.add_paragraph()
                    p.add_run("Notes: ").bold = True
                    p.add_run(record.notes)
                
                p = doc.add_paragraph()
                p.add_run("Source Files: ").bold = True
                p.add_run(', '.join(record.source_files))
                
                p = doc.add_paragraph()
                p.add_run("Confidence: ").bold = True
                p.add_run(f"{record.confidence_score:.2f}")
                
                doc.add_page_break()
            
            doc.save(docx_file)
            logger.info(f"Generated medical DOCX: {docx_file}")
            return [str(docx_file)]
            
        except Exception as e:
            logger.error(f"Failed to write medical DOCX: {e}")
            return []
    
    def write_life_docx(self, data: ConsolidatedData) -> List[str]:
        """Write life history data to DOCX format."""
        if not DOCX_AVAILABLE:
            logger.warning("python-docx not available, skipping DOCX generation")
            return []
        
        docx_file = self.docx_dir / f"life_history_{self.timestamp}.docx"
        
        try:
            doc = Document()
            
            # Title
            title = doc.add_heading('Life History Summary', 0)
            title.alignment = WD_ALIGN_PARAGRAPH.CENTER
            
            # Metadata
            doc.add_paragraph(f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
            doc.add_paragraph(f"Total Records: {len(data.life_history_records)}")
            doc.add_paragraph(f"Source Files: {len(data.source_files)}")
            
            # Sort records by start date
            sorted_records = sorted(data.life_history_records, 
                                  key=lambda r: r.start_date or datetime.min)
            
            # Records
            for i, record in enumerate(sorted_records, 1):
                doc.add_heading(f"{i}. {record.role}", level=1)
                
                # Dates
                start_str = record.start_date.strftime('%Y-%m-%d') if record.start_date else 'Unknown'
                end_str = record.end_date.strftime('%Y-%m-%d') if record.end_date else 'Ongoing'
                p = doc.add_paragraph()
                p.add_run("Period: ").bold = True
                p.add_run(f"{start_str} to {end_str}")
                
                if record.location:
                    p = doc.add_paragraph()
                    p.add_run("Location: ").bold = True
                    p.add_run(record.location)
                
                if record.key_achievements:
                    p = doc.add_paragraph()
                    p.add_run("Key Achievements:").bold = True
                    for achievement in record.key_achievements:
                        doc.add_paragraph(f"• {achievement}", style='List Bullet')
                
                p = doc.add_paragraph()
                p.add_run("Source Files: ").bold = True
                p.add_run(', '.join(record.source_files))
                
                p = doc.add_paragraph()
                p.add_run("Confidence: ").bold = True
                p.add_run(f"{record.confidence_score:.2f}")
                
                doc.add_page_break()
            
            doc.save(docx_file)
            logger.info(f"Generated life history DOCX: {docx_file}")
            return [str(docx_file)]
            
        except Exception as e:
            logger.error(f"Failed to write life history DOCX: {e}")
            return []
    
    def backup_json_responses(self, responses: List[GPTResponse], mode: str) -> List[str]:
        """Backup raw GPT responses to JSON."""
        json_file = self.json_dir / f"raw_responses_{mode}_{self.timestamp}.json"
        
        try:
            backup_data = {
                "timestamp": datetime.now().isoformat(),
                "mode": mode,
                "total_responses": len(responses),
                "responses": []
            }
            
            for response in responses:
                backup_data["responses"].append({
                    "chunk_id": response.chunk_id,
                    "structured_data": response.structured_data,
                    "summary": response.summary,
                    "confidence_score": response.confidence_score,
                    "source_references": response.source_references,
                    "processing_time": response.processing_time,
                    "token_usage": response.token_usage,
                    "raw_response": response.raw_response
                })
            
            with open(json_file, 'w', encoding='utf-8') as f:
                json.dump(backup_data, f, indent=2, ensure_ascii=False)
            
            logger.info(f"Backed up {len(responses)} responses to: {json_file}")
            return [str(json_file)]
            
        except Exception as e:
            logger.error(f"Failed to backup JSON responses: {e}")
            return []
    
    def write_processing_metadata(self, data: ConsolidatedData, mode: str) -> List[str]:
        """Write processing metadata to JSON."""
        metadata_file = self.json_dir / f"processing_metadata_{mode}_{self.timestamp}.json"
        
        try:
            metadata = {
                "timestamp": datetime.now().isoformat(),
                "mode": mode,
                "processing_metadata": data.processing_metadata,
                "deduplication_summary": {
                    "total_merges": len(data.deduplication_log),
                    "merge_actions": data.deduplication_log
                },
                "output_summary": {
                    "medical_records": len(data.medical_records),
                    "life_history_records": len(data.life_history_records),
                    "source_files": len(data.source_files),
                    "summaries": len(data.summaries)
                }
            }
            
            with open(metadata_file, 'w', encoding='utf-8') as f:
                json.dump(metadata, f, indent=2, ensure_ascii=False)
            
            logger.info(f"Generated processing metadata: {metadata_file}")
            return [str(metadata_file)]
            
        except Exception as e:
            logger.error(f"Failed to write processing metadata: {e}")
            return []

    def write_code_markdown(self, data: ConsolidatedData) -> List[str]:
        """Write code review analysis to Markdown format."""
        md_file = self.md_dir / f"code_review_{self.timestamp}.md"
        md_file.parent.mkdir(parents=True, exist_ok=True)
        
        with open(md_file, 'w', encoding='utf-8') as f:
            f.write(f"# Code Review Analysis Report\n\n")
            f.write(f"**Generated:** {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}\n")
            f.write(f"**Total Files Reviewed:** {len(data.code_review_records)}\n")
            f.write(f"**Source Files:** {len(data.source_files)}\n\n")
            
            # Individual File Reviews
            f.write(f"## File Reviews\n\n")
            
            for i, record in enumerate(data.code_review_records, 1):
                f.write(f"### {i}. {record.file_name}\n\n")
                f.write(f"**Language:** {record.language}\n")
                f.write(f"**Overall Quality:** {record.overall_quality}\n")
                f.write(f"**Quality Score:** {record.quality_score}/10\n")
                f.write(f"**Improvement Priority:** {record.improvement_priority}\n")
                f.write(f"**Estimated Effort:** {record.estimated_effort}\n\n")
                
                # Security Issues
                if record.security_issues:
                    f.write(f"**🔒 Security Issues:**\n")
                    for issue in record.security_issues:
                        f.write(f"- **{issue.get('severity', 'Unknown')}:** {issue.get('issue', 'N/A')}\n")
                        f.write(f"  - *Location:* {issue.get('line_reference', 'N/A')}\n")
                        f.write(f"  - *Fix:* {issue.get('recommendation', 'N/A')}\n")
                    f.write(f"\n")
                
                # Performance Issues
                if record.performance_issues:
                    f.write(f"**⚡ Performance Issues:**\n")
                    for issue in record.performance_issues:
                        f.write(f"- **{issue.get('severity', 'Unknown')}:** {issue.get('issue', 'N/A')}\n")
                        f.write(f"  - *Location:* {issue.get('line_reference', 'N/A')}\n")
                        f.write(f"  - *Fix:* {issue.get('recommendation', 'N/A')}\n")
                    f.write(f"\n")
                
                # Positive Aspects
                if record.positive_aspects:
                    f.write(f"**✅ Positive Aspects:**\n")
                    for aspect in record.positive_aspects:
                        f.write(f"- {aspect}\n")
                    f.write(f"\n")
                
                f.write(f"**Confidence:** {record.confidence:.2f}\n\n")
                f.write(f"---\n\n")
        
        logger.info(f"Generated code review markdown: {md_file}")
        return [str(md_file)]

    def write_code_csv(self, data: ConsolidatedData) -> List[str]:
        """Write code review data to CSV format."""
        csv_file = self.csv_dir / f"code_review_{self.timestamp}.csv"
        csv_file.parent.mkdir(parents=True, exist_ok=True)
        
        with open(csv_file, 'w', newline='', encoding='utf-8') as f:
            writer = csv.writer(f)
            
            # Write header
            writer.writerow([
                'File Name', 'Language', 'Overall Quality', 'Quality Score',
                'Security Issues Count', 'Performance Issues Count', 
                'Architecture Suggestions Count', 'Positive Aspects Count',
                'Improvement Priority', 'Estimated Effort', 'Confidence'
            ])
            
            # Write data rows
            for record in data.code_review_records:
                writer.writerow([
                    record.file_name,
                    record.language,
                    record.overall_quality,
                    record.quality_score,
                    len(record.security_issues),
                    len(record.performance_issues),
                    len(record.architecture_suggestions),
                    len(record.positive_aspects),
                    record.improvement_priority,
                    record.estimated_effort,
                    record.confidence
                ])
        
        logger.info(f"Generated code review CSV: {csv_file}")
        return [str(csv_file)]

    def write_code_docx(self, data: ConsolidatedData) -> List[str]:
        """Write code review analysis to DOCX format."""
        if not DOCX_AVAILABLE:
            logger.warning("python-docx not available, skipping DOCX generation")
            return []
        
        docx_file = self.docx_dir / f"code_review_{self.timestamp}.docx"
        docx_file.parent.mkdir(parents=True, exist_ok=True)
        
        doc = Document()
        
        # Title
        title = doc.add_heading('Code Review Analysis Report', 0)
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        
        # File Reviews
        doc.add_heading('File Reviews', level=1)
        
        for i, record in enumerate(data.code_review_records, 1):
            doc.add_heading(f'{i}. {record.file_name}', level=2)
            
            # Basic info
            doc.add_paragraph(f"Language: {record.language}")
            doc.add_paragraph(f"Quality Score: {record.quality_score}/10")
            doc.add_paragraph(f"Priority: {record.improvement_priority}")
            
            # Positive aspects
            if record.positive_aspects:
                doc.add_heading('Positive Aspects', level=3)
                for aspect in record.positive_aspects:
                    doc.add_paragraph(aspect, style='List Bullet')
        
        doc.save(str(docx_file))
        logger.info(f"Generated code review DOCX: {docx_file}")
        return [str(docx_file)]


def is_docx_available() -> bool:
    """Check if python-docx is available."""
    return DOCX_AVAILABLE