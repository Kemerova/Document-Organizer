"""
File loading utilities for Document Organizer.
Handles discovery and content extraction from multiple file formats.
"""

import os
import logging
from pathlib import Path
from typing import List, Dict, Optional, Generator
import chardet
import docx
import PyPDF2
from PIL import Image
import io

from .models import DocumentContent

logger = logging.getLogger(__name__)

# Supported file extensions
SUPPORTED_EXTENSIONS = {'.txt', '.docx', '.pdf', '.jpg', '.jpeg', '.png'}
TEXT_EXTENSIONS = {'.txt'}
DOCX_EXTENSIONS = {'.docx'}
PDF_EXTENSIONS = {'.pdf'}
IMAGE_EXTENSIONS = {'.jpg', '.jpeg', '.png'}


class FileLoadError(Exception):
    """Raised when file loading fails."""
    pass


class FileLoader:
    """Handles file discovery and content loading."""
    
    def __init__(self):
        """Initialize file loader."""
        self.supported_extensions = SUPPORTED_EXTENSIONS
    
    def discover_files(self, input_path: str, recursive: bool = True) -> List[str]:
        """
        Discover all supported files in the given path.
        
        Args:
            input_path: Path to search for files
            recursive: Whether to search subdirectories
            
        Returns:
            List of file paths
            
        Raises:
            FileLoadError: If input path doesn't exist
        """
        input_path = Path(input_path)
        
        if not input_path.exists():
            raise FileLoadError(f"Input path does not exist: {input_path}")
        
        files = []
        
        if input_path.is_file():
            if input_path.suffix.lower() in self.supported_extensions:
                files.append(str(input_path))
            else:
                logger.warning(f"Unsupported file type: {input_path}")
        else:
            # Directory search
            pattern = "**/*" if recursive else "*"
            for file_path in input_path.glob(pattern):
                if file_path.is_file() and file_path.suffix.lower() in self.supported_extensions:
                    files.append(str(file_path))
        
        logger.info(f"Discovered {len(files)} supported files in {input_path}")
        return sorted(files)
    
    def load_file(self, file_path: str) -> DocumentContent:
        """
        Load content from a file based on its extension.
        
        Args:
            file_path: Path to the file to load
            
        Returns:
            DocumentContent object with extracted content
            
        Raises:
            FileLoadError: If file loading fails
        """
        file_path = Path(file_path)
        
        if not file_path.exists():
            raise FileLoadError(f"File does not exist: {file_path}")
        
        extension = file_path.suffix.lower()
        
        try:
            if extension in TEXT_EXTENSIONS:
                return self.load_text_file(str(file_path))
            elif extension in DOCX_EXTENSIONS:
                return self.load_docx_file(str(file_path))
            elif extension in PDF_EXTENSIONS:
                return self.load_pdf_file(str(file_path))
            elif extension in IMAGE_EXTENSIONS:
                # Image files will be handled by OCR module
                return DocumentContent(
                    filename=str(file_path),
                    content="",  # Will be filled by OCR
                    file_type=extension[1:],  # Remove the dot
                    ocr_applied=False,
                    extraction_confidence=0.0
                )
            else:
                raise FileLoadError(f"Unsupported file type: {extension}")
                
        except Exception as e:
            logger.error(f"Error loading file {file_path}: {e}")
            raise FileLoadError(f"Failed to load {file_path}: {e}")
    
    def load_text_file(self, file_path: str) -> DocumentContent:
        """
        Load content from a text file with encoding detection.
        
        Args:
            file_path: Path to the text file
            
        Returns:
            DocumentContent object
        """
        file_path = Path(file_path)
        
        # Detect encoding
        with open(file_path, 'rb') as f:
            raw_data = f.read()
            encoding_result = chardet.detect(raw_data)
            encoding = encoding_result.get('encoding', 'utf-8')
            confidence = encoding_result.get('confidence', 0.0)
        
        # Read file with detected encoding
        try:
            with open(file_path, 'r', encoding=encoding) as f:
                content = f.read()
        except UnicodeDecodeError:
            # Fallback to utf-8 with error handling
            logger.warning(f"Encoding detection failed for {file_path}, using utf-8 with error handling")
            with open(file_path, 'r', encoding='utf-8', errors='replace') as f:
                content = f.read()
            confidence = 0.5
        
        return DocumentContent(
            filename=str(file_path),
            content=content,
            page_numbers=[1],  # Text files are considered single page
            file_type="txt",
            ocr_applied=False,
            extraction_confidence=confidence
        )
    
    def load_docx_file(self, file_path: str) -> DocumentContent:
        """
        Load content from a DOCX file.
        
        Args:
            file_path: Path to the DOCX file
            
        Returns:
            DocumentContent object
        """
        try:
            doc = docx.Document(file_path)
            
            # Extract text from all paragraphs
            paragraphs = []
            for paragraph in doc.paragraphs:
                if paragraph.text.strip():
                    paragraphs.append(paragraph.text.strip())
            
            # Extract text from tables
            for table in doc.tables:
                for row in table.rows:
                    row_text = []
                    for cell in row.cells:
                        if cell.text.strip():
                            row_text.append(cell.text.strip())
                    if row_text:
                        paragraphs.append(" | ".join(row_text))
            
            content = "\\n\\n".join(paragraphs)
            
            # Estimate page count (rough approximation)
            estimated_pages = max(1, len(content) // 2000)  # ~2000 chars per page
            page_numbers = list(range(1, estimated_pages + 1))
            
            return DocumentContent(
                filename=str(file_path),
                content=content,
                page_numbers=page_numbers,
                file_type="docx",
                ocr_applied=False,
                extraction_confidence=0.95  # DOCX extraction is usually reliable
            )
            
        except Exception as e:
            logger.error(f"Error reading DOCX file {file_path}: {e}")
            raise FileLoadError(f"Failed to read DOCX file: {e}")
    
    def load_pdf_file(self, file_path: str) -> DocumentContent:
        """
        Load content from a PDF file with page tracking.
        
        Args:
            file_path: Path to the PDF file
            
        Returns:
            DocumentContent object
        """
        try:
            content_parts = []
            page_numbers = []
            
            with open(file_path, 'rb') as file:
                pdf_reader = PyPDF2.PdfReader(file)
                
                for page_num, page in enumerate(pdf_reader.pages, 1):
                    try:
                        page_text = page.extract_text()
                        if page_text.strip():
                            content_parts.append(f"[Page {page_num}]\\n{page_text.strip()}")
                            page_numbers.append(page_num)
                        else:
                            # Empty page or might need OCR
                            logger.debug(f"Page {page_num} in {file_path} appears empty or may need OCR")
                            page_numbers.append(page_num)
                    except Exception as e:
                        logger.warning(f"Error extracting text from page {page_num} in {file_path}: {e}")
                        page_numbers.append(page_num)
            
            content = "\\n\\n".join(content_parts)
            
            # Determine if this might be a scanned PDF (low text content)
            is_likely_scanned = len(content.strip()) < 100 and len(page_numbers) > 0
            confidence = 0.3 if is_likely_scanned else 0.9
            
            return DocumentContent(
                filename=str(file_path),
                content=content,
                page_numbers=page_numbers,
                file_type="pdf",
                ocr_applied=False,
                extraction_confidence=confidence
            )
            
        except Exception as e:
            logger.error(f"Error reading PDF file {file_path}: {e}")
            raise FileLoadError(f"Failed to read PDF file: {e}")
    
    def get_file_info(self, file_path: str) -> Dict[str, any]:
        """
        Get basic information about a file.
        
        Args:
            file_path: Path to the file
            
        Returns:
            Dictionary with file information
        """
        file_path = Path(file_path)
        
        if not file_path.exists():
            return {}
        
        stat = file_path.stat()
        
        return {
            'filename': file_path.name,
            'full_path': str(file_path),
            'extension': file_path.suffix.lower(),
            'size_bytes': stat.st_size,
            'size_mb': round(stat.st_size / (1024 * 1024), 2),
            'modified_time': stat.st_mtime,
            'is_supported': file_path.suffix.lower() in self.supported_extensions
        }
    
    def batch_load_files(self, file_paths: List[str]) -> Generator[DocumentContent, None, None]:
        """
        Load multiple files in batch, yielding results as they're processed.
        
        Args:
            file_paths: List of file paths to load
            
        Yields:
            DocumentContent objects for successfully loaded files
        """
        for file_path in file_paths:
            try:
                yield self.load_file(file_path)
            except FileLoadError as e:
                logger.error(f"Failed to load {file_path}: {e}")
                continue
    
    def validate_input_path(self, input_path: str) -> bool:
        """
        Validate that the input path exists and contains supported files.
        
        Args:
            input_path: Path to validate
            
        Returns:
            True if path is valid and contains supported files
        """
        try:
            files = self.discover_files(input_path)
            return len(files) > 0
        except FileLoadError:
            return False
    
    def get_file_statistics(self, file_paths: List[str]) -> Dict[str, any]:
        """
        Get statistics about a collection of files.
        
        Args:
            file_paths: List of file paths
            
        Returns:
            Dictionary with file statistics
        """
        stats = {
            'total_files': len(file_paths),
            'by_extension': {},
            'total_size_mb': 0,
            'largest_file': None,
            'smallest_file': None
        }
        
        file_sizes = []
        
        for file_path in file_paths:
            info = self.get_file_info(file_path)
            if not info:
                continue
                
            # Count by extension
            ext = info['extension']
            stats['by_extension'][ext] = stats['by_extension'].get(ext, 0) + 1
            
            # Track sizes
            size_mb = info['size_mb']
            stats['total_size_mb'] += size_mb
            file_sizes.append((file_path, size_mb))
        
        # Find largest and smallest files
        if file_sizes:
            file_sizes.sort(key=lambda x: x[1])
            stats['smallest_file'] = file_sizes[0]
            stats['largest_file'] = file_sizes[-1]
        
        return stats